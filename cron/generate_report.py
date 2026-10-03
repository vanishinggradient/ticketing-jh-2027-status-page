"""
generate_report.py
==================
Generates comparison.xlsx with JH 2027 vs JH 2026 sales pace analysis.
Maintains history.json (one entry per day). Outputs comparison.xlsx.

Runs after monitor.py in the GitHub Actions workflow.
No API calls. Reads status.json, writes history.json + comparison.xlsx.
"""

import json
import pathlib
from datetime import date, datetime, timezone, timedelta

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment

# ── Constants ─────────────────────────────────────────────────────────────────

EVENT_DATE = date(2027, 1, 9)
IST = timezone(timedelta(hours=5, minutes=30))
REPO_ROOT = pathlib.Path(__file__).resolve().parent.parent

# JH 2026 historical data from Nivi's spreadsheet.
# Tuple: (days_to_jan9, cumulative_sold, date_label, note)
JH_2026 = [
    (63,  0,     "Nov 7",  ""),
    (62,  1000,  "Nov 8",  "Early Bird Sold Out"),
    (59,  1200,  "Nov 11", ""),
    (49,  1500,  "Nov 21", ""),
    (36,  2000,  "Dec 4",  ""),
    (18,  2358,  "Dec 22", ""),
    (17,  2412,  "Dec 23", ""),
    (16,  2465,  "Dec 24", ""),
    (15,  2689,  "Dec 25", "Skibidi uncles"),
    (14,  3201,  "Dec 26", ""),
    (13,  3447,  "Dec 27", ""),
    (12,  3675,  "Dec 28", ""),
    (11,  3884,  "Dec 29", ""),
    (10,  4054,  "Dec 30", ""),
    (9,   4309,  "Dec 31", ""),
    (8,   4620,  "Jan 1",  ""),
    (7,   5002,  "Jan 2",  ""),
    (6,   5300,  "Jan 3",  ""),
    (5,   5663,  "Jan 4",  "Weeb Central Post"),
    (4,   6277,  "Jan 5",  ""),
    (3,   6805,  "Jan 6",  "Gurukiran reel, JP interview"),
    (2,   7309,  "Jan 7",  ""),
    (1,   8059,  "Jan 8",  ""),
    (0,   9850,  "Jan 9",  "Price raised to Rs. 750 at night"),
    (-1,  11500, "Jan 10", ""),
    (-2,  12100, "Jan 11", "Sold out"),
]

# Pace milestones for JH 2027.
# (label, milestone_date, days_to_event, target_sold)
PACE_MILESTONES = [
    ("Nov 1, 2026",  date(2026, 11, 1),  69,  3000),
    ("Dec 1, 2026",  date(2026, 12, 1),  39,  6500),
    ("Jan 1, 2027",  date(2027, 1, 1),   8,   12000),
    ("Jan 9, 2027",  date(2027, 1, 9),   0,   18000),
]


# ── Helpers ───────────────────────────────────────────────────────────────────

def today_ist() -> date:
    return datetime.now(IST).date()


def days_remaining(from_date: date = None) -> int:
    return (EVENT_DATE - (from_date or today_ist())).days


def interpolate_target(d: int) -> int | None:
    """Linear interpolation between milestone targets, given days_to_event."""
    anchors = [(100, 0), (69, 3000), (39, 6500), (8, 12000), (0, 18000)]
    for i in range(len(anchors) - 1):
        d1, t1 = anchors[i]
        d2, t2 = anchors[i + 1]
        if d2 <= d <= d1:
            frac = (d1 - d) / (d1 - d2)
            return round(t1 + frac * (t2 - t1))
    return 18000 if d < 0 else None


def load_json(path: pathlib.Path, default):
    if path.exists():
        with open(path) as f:
            return json.load(f)
    return default


def save_json(path: pathlib.Path, data):
    with open(path, "w") as f:
        json.dump(data, f, indent=2)


# ── Style helpers ─────────────────────────────────────────────────────────────

def hdr_cell(ws, row: int, col: int, value, bg: str = "312E81"):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(bold=True, color="FFFFFF", size=10)
    c.fill = PatternFill("solid", fgColor=bg)
    c.alignment = Alignment(horizontal="center", vertical="center")
    return c


def set_col_widths(ws, widths: list[tuple[str, float]]):
    for col_letter, width in widths:
        ws.column_dimensions[col_letter].width = width


def rag_fill(pct: float) -> PatternFill:
    if pct >= 0.9:
        return PatternFill("solid", fgColor="D1FAE5")  # green
    if pct >= 0.7:
        return PatternFill("solid", fgColor="FEF3C7")  # amber
    return PatternFill("solid", fgColor="FEE2E2")      # red


def rag_label(pct: float) -> str:
    if pct >= 0.9:
        return "On Track"
    if pct >= 0.7:
        return "At Risk"
    return "Off Track"


# ── Sheet builders ────────────────────────────────────────────────────────────

def build_pace_chart_sheet(wb: Workbook, history: list):
    """
    Sheet 1: Pace Chart
    Unified day-by-day table (D-100 to D+2), 3 series:
      - JH 2026 (navy)
      - JH 2027 (orange, from history.json)
      - Target pace line (red dashed)
    Embedded LineChart.
    """
    ws = wb.active
    ws.title = "Pace Chart"
    ws.row_dimensions[1].height = 18

    for col, title in enumerate(
        ["Days to Jan 9", "JH 2026 (sold)", "JH 2027 (sold)", "Target (pace)"], 1
    ):
        hdr_cell(ws, 1, col, title)

    jh2026_lookup = {d: s for d, s, _, _ in JH_2026}
    jh2027_lookup = {e["days_to_event"]: e["total_sold"] for e in history}

    # One row per day from D-100 down to D+2
    all_days = list(range(100, -3, -1))
    for i, d in enumerate(all_days, 2):
        ws.cell(row=i, column=1, value=d).font = Font(size=9, color="6B7280")
        if d in jh2026_lookup:
            ws.cell(row=i, column=2, value=jh2026_lookup[d]).font = Font(size=9, color="1E3A5F")
        if d in jh2027_lookup:
            ws.cell(row=i, column=3, value=jh2027_lookup[d]).font = Font(size=9, bold=True, color="C2410C")
        t = interpolate_target(d)
        if t is not None:
            ws.cell(row=i, column=4, value=t).font = Font(size=9, color="991B1B")

    set_col_widths(ws, [("A", 14), ("B", 16), ("C", 16), ("D", 16)])

    n = len(all_days)

    # Line chart
    chart = LineChart()
    chart.title = "Sales Pace: JH 2027 vs JH 2026"
    chart.style = 10
    chart.height = 14
    chart.width = 24
    chart.y_axis.title = "Tickets Sold"
    chart.x_axis.title = "Days to Jan 9 (left = far, right = near)"

    cats = Reference(ws, min_col=1, min_row=2, max_row=n + 1)

    for col in (2, 3, 4):
        data = Reference(ws, min_col=col, min_row=1, max_row=n + 1)
        chart.add_data(data, titles_from_data=True)

    chart.set_categories(cats)

    # JH 2026: navy solid
    chart.series[0].graphicalProperties.line.solidFill = "1E3A5F"
    chart.series[0].graphicalProperties.line.width = 20000

    # JH 2027: orange solid, thicker
    chart.series[1].graphicalProperties.line.solidFill = "F97316"
    chart.series[1].graphicalProperties.line.width = 25000

    # Target: red, thinner
    chart.series[2].graphicalProperties.line.solidFill = "DC2626"
    chart.series[2].graphicalProperties.line.width = 14000
    chart.series[2].graphicalProperties.line.dashDot = "dash"

    ws.add_chart(chart, "F2")


def build_targets_sheet(wb: Workbook, history: list):
    """
    Sheet 2: Pace Targets
    RAG status table: milestone, target, actual, delta, status.
    """
    ws = wb.create_sheet("Pace Targets")
    ws.row_dimensions[1].height = 18

    for col, h in enumerate(
        ["Milestone", "Date", "Days to Event", "Target", "Actual", "Delta", "Status"], 1
    ):
        hdr_cell(ws, 1, col, h, bg="9D174D")

    today = today_ist()
    sold_by_date = {e["date"]: e["total_sold"] for e in history}
    current_sold = history[-1]["total_sold"] if history else 0

    for row, (label, m_date, dte, target) in enumerate(PACE_MILESTONES, 2):
        ws.cell(row=row, column=1, value=label).font = Font(bold=True, size=10)
        ws.cell(row=row, column=2, value=m_date.strftime("%b %d, %Y")).font = Font(size=10)
        ws.cell(row=row, column=3, value=dte).font = Font(size=10, color="6B7280")
        ws.cell(row=row, column=4, value=target).font = Font(size=10)

        if m_date > today:
            # Not yet reached — show current total as "running total so far"
            ws.cell(row=row, column=5, value=current_sold).font = Font(size=10, color="6B7280", italic=True)
            ws.cell(row=row, column=6, value="").font = Font(size=10)
            status_cell = ws.cell(row=row, column=7, value="Upcoming")
            status_cell.font = Font(size=10, color="6B7280")
            status_cell.fill = PatternFill("solid", fgColor="F3F4F6")
        else:
            # Past milestone — find nearest actual
            actual = sold_by_date.get(m_date.isoformat())
            if actual is None and history:
                # Nearest by days_to_event
                nearest = min(history, key=lambda e: abs(e["days_to_event"] - dte))
                actual = nearest["total_sold"]

            if actual is not None:
                delta = actual - target
                pct = actual / target if target > 0 else 0
                ws.cell(row=row, column=5, value=actual).font = Font(bold=True, size=10)
                delta_cell = ws.cell(row=row, column=6, value=delta)
                delta_cell.font = Font(size=10, color="16A34A" if delta >= 0 else "DC2626", bold=True)
                status_cell = ws.cell(row=row, column=7, value=rag_label(pct))
                status_cell.font = Font(bold=True, size=10)
                status_cell.fill = rag_fill(pct)
            else:
                ws.cell(row=row, column=5, value="No data").font = Font(size=10, color="9CA3AF", italic=True)

    set_col_widths(ws, [("A", 18), ("B", 16), ("C", 14), ("D", 12), ("E", 12), ("F", 10), ("G", 12)])


def build_breakdown_sheet(wb: Workbook, status: dict):
    """
    Sheet 3: Live Breakdown
    Current ticket state from status.json.
    """
    ws = wb.create_sheet("Live Breakdown")

    ts = status.get("updated_at", "unknown")
    ws.cell(row=1, column=1, value=f"Last updated: {ts}").font = Font(italic=True, color="6B7280", size=9)
    ws.merge_cells("A1:F1")

    for col, h in enumerate(["Ticket", "Price (Rs.)", "Sold", "Capacity", "Remaining", "% Full"], 1):
        hdr_cell(ws, 2, col, h)

    tickets = status.get("tickets", [])
    for row, t in enumerate(tickets, 3):
        cap = t.get("capacity", 0) or 0
        sold = t.get("sold", 0) or 0
        remaining = t.get("remaining", 0) or 0
        pct = round(sold / cap * 100, 1) if cap > 0 else 0
        soldout = remaining == 0

        row_font = Font(size=10, color="9CA3AF" if soldout else "1F2937")
        for col, val in enumerate([t.get("name", ""), t.get("price", 0), sold, cap, remaining, pct], 1):
            c = ws.cell(row=row, column=col, value=val)
            c.font = row_font
            if soldout:
                c.fill = PatternFill("solid", fgColor="F9FAFB")

        # % Full: colour-coded
        pct_cell = ws.cell(row=row, column=6)
        if not soldout:
            if pct >= 80:
                pct_cell.font = Font(size=10, bold=True, color="DC2626")
            elif pct >= 50:
                pct_cell.font = Font(size=10, color="D97706")
            else:
                pct_cell.font = Font(size=10, color="16A34A")

    set_col_widths(ws, [("A", 42), ("B", 12), ("C", 10), ("D", 10), ("E", 12), ("F", 10)])


def build_jh2026_raw_sheet(wb: Workbook):
    """
    Sheet 4: JH 2026 Raw
    The original spreadsheet data for reference.
    """
    ws = wb.create_sheet("JH 2026 Raw")
    ws.row_dimensions[1].height = 18

    for col, h in enumerate(["Date", "Days to Jan 9", "Cumulative Sold", "Daily Change", "Note"], 1):
        hdr_cell(ws, 1, col, h, bg="1E3A5F")

    prev_sold = 0
    for row, (d, sold, date_label, note) in enumerate(JH_2026, 2):
        delta = sold - prev_sold
        prev_sold = sold

        ws.cell(row=row, column=1, value=date_label).font = Font(size=10)
        ws.cell(row=row, column=2, value=d).font = Font(size=10, color="6B7280")
        ws.cell(row=row, column=3, value=sold).font = Font(size=10, bold=True)

        delta_cell = ws.cell(row=row, column=4, value=f"+{delta:,}" if delta > 0 else "")
        delta_cell.font = Font(size=10, color="16A34A" if delta > 500 else "6B7280")

        ws.cell(row=row, column=5, value=note).font = Font(size=10, italic=True, color="6B7280")

    # Highlight big spikes
    for row, (d, sold, _, note) in enumerate(JH_2026, 2):
        if note:
            for col in range(1, 6):
                ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor="FEF9C3")

    set_col_widths(ws, [("A", 10), ("B", 14), ("C", 16), ("D", 14), ("E", 32)])


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    print("generate_report.py starting...")

    # Load status.json
    status = load_json(REPO_ROOT / "status.json", {"tickets": [], "updated_at": "unknown"})
    total_sold = sum(t.get("sold", 0) for t in status.get("tickets", []))
    dte = days_remaining()
    today = today_ist().isoformat()

    print(f"  Total sold today: {total_sold}")
    print(f"  Days to event: {dte}")

    # Update history.json
    history = load_json(REPO_ROOT / "history.json", [])
    updated = False
    for entry in history:
        if entry["date"] == today:
            entry["total_sold"] = total_sold
            entry["days_to_event"] = dte
            updated = True
            break
    if not updated:
        history.append({"date": today, "days_to_event": dte, "total_sold": total_sold})
        history.sort(key=lambda x: x["date"])

    save_json(REPO_ROOT / "history.json", history)
    print(f"  History updated: {len(history)} entries")

    # Build workbook
    wb = Workbook()
    build_pace_chart_sheet(wb, history)
    build_targets_sheet(wb, history)
    build_breakdown_sheet(wb, status)
    build_jh2026_raw_sheet(wb)

    out = REPO_ROOT / "comparison.xlsx"
    wb.save(out)
    print(f"  Saved: {out}")
    print("Done.")


if __name__ == "__main__":
    main()
